"""Resume text extraction (PDF / DOCX / TXT) and section segmentation."""
from __future__ import annotations

import io
import re
import unicodedata

SECTION_ALIASES: dict[str, list[str]] = {
    "summary": ["summary", "professional summary", "profile", "about me", "objective", "career objective",
                "professional profile", "career summary", "overview"],
    "experience": ["experience", "work experience", "professional experience", "employment history",
                   "work history", "career history", "employment", "relevant experience"],
    "education": ["education", "academic background", "academics", "education and training", "qualifications",
                  "academic qualifications"],
    "skills": ["skills", "technical skills", "core competencies", "key skills", "competencies", "technologies",
               "tech stack", "skills and tools", "tools", "expertise", "areas of expertise"],
    "projects": ["projects", "personal projects", "key projects", "academic projects", "selected projects"],
    "certifications": ["certifications", "certificates", "licenses", "licenses and certifications",
                       "certifications and licenses", "courses"],
    "awards": ["awards", "honors", "achievements", "honors and awards", "accomplishments"],
    "publications": ["publications", "research"],
    "languages": ["languages"],
    "volunteer": ["volunteer", "volunteering", "volunteer experience"],
    "interests": ["interests", "hobbies"],
}
_HEADING_LOOKUP = {alias: key for key, aliases in SECTION_ALIASES.items() for alias in aliases}

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s.-]?)?(?:\(?\d{2,5}\)?[\s.-]?)?\d{3,5}[\s.-]?\d{3,5}")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[\w-]+/?", re.I)
GITHUB_RE = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[\w-]+/?", re.I)
URL_RE = re.compile(r"https?://\S+")


class ParseError(ValueError):
    pass


def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = _pdf_text(data)
    elif name.endswith(".docx"):
        text = _docx_text(data)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ParseError("Unsupported file type. Please upload PDF, DOCX or TXT.")
    text = normalize_text(text)
    if len(text.strip()) < 30:
        raise ParseError("Could not extract readable text. If this is a scanned PDF, please upload a text-based file.")
    return text


def _pdf_text(data: bytes) -> str:
    import pdfplumber

    pages = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text(x_tolerance=1.5, y_tolerance=3) or "")
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Invalid PDF file: {exc}") from exc
    return "\n".join(pages)


def _docx_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Invalid DOCX file: {exc}") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    return "\n".join(lines)


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = re.sub(r"[\u2022\u25cf\u25aa\u25a0\u2023\u2043\u2219\uf0b7\uf0a7]", "\u2022", text)
    text = re.sub(r"[ \u00a0]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _heading_key(line: str) -> str | None:
    clean = re.sub(r"[^a-zA-Z& ]", "", line).strip().lower().replace("&", "and")
    clean = re.sub(r"\s+", " ", clean)
    if not clean or len(clean.split()) > 5:
        return None
    return _HEADING_LOOKUP.get(clean)


def segment_sections(text: str) -> dict[str, str]:
    """Split resume text into canonical sections. Text before the first heading is the 'header'."""
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.split("\n"):
        line = raw.strip()
        key = _heading_key(line) if line and len(line) < 60 else None
        if key:
            current = key
            sections.setdefault(current, [])
            continue
        # Support "Skills: Python, SQL" inline headings
        m = re.match(r"^([A-Za-z ]{3,30}):\s*(.+)$", line)
        if m and _heading_key(m.group(1)):
            k = _heading_key(m.group(1))
            sections.setdefault(k, []).append(m.group(2))
            continue
        sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items() if "\n".join(v).strip()}


def extract_contact(text: str, header: str = "") -> dict:
    head = header or text[:600]
    email = EMAIL_RE.search(text)
    phone = None
    for m in PHONE_RE.finditer(head + "\n" + text[:1500]):
        digits = re.sub(r"\D", "", m.group())
        if 9 <= len(digits) <= 15 and not re.fullmatch(r"(19|20)\d{2}(19|20)\d{2}", digits):
            phone = m.group().strip()
            break
    linkedin = LINKEDIN_RE.search(text)
    github = GITHUB_RE.search(text)
    return {
        "name": guess_name(head),
        "email": email.group() if email else None,
        "phone": phone,
        "linkedin": linkedin.group() if linkedin else None,
        "github": github.group() if github else None,
    }


def guess_name(header: str) -> str | None:
    for line in header.split("\n")[:5]:
        line = line.strip()
        if not line or EMAIL_RE.search(line) or any(ch.isdigit() for ch in line) or "http" in line.lower():
            continue
        words = re.split(r"\s+", re.sub(r"[|,•].*$", "", line).strip())
        if 1 < len(words) <= 4 and all(w[:1].isupper() for w in words if w):
            return " ".join(words)
    try:
        from app.services.nlp import spacy_person

        return spacy_person(header)
    except Exception:  # noqa: BLE001
        return None


def parse_resume(data: bytes, filename: str) -> tuple[str, dict[str, str], dict]:
    text = extract_text(data, filename)
    sections = segment_sections(text)
    contact = extract_contact(text, sections.get("header", ""))
    return text, sections, contact
