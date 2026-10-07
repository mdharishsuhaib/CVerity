import io

import docx

from app.services.parser import extract_text, parse_resume, segment_sections


def test_sections_and_contact(backend_resume):
    text, sections, contact = parse_resume(backend_resume, "r.txt")
    assert {"summary", "experience", "education", "skills", "certifications"} <= set(sections)
    assert contact["email"] == "priya.sharma@example.com"
    assert contact["name"] == "Priya Sharma"
    assert contact["linkedin"] and contact["github"] and contact["phone"]
    assert contact["phone"].replace(" ", "").endswith("43210")


def test_docx_extraction():
    d = docx.Document()
    d.add_paragraph("Jane Doe")
    d.add_paragraph("Skills")
    d.add_paragraph("Python, SQL, Tableau and more relevant skills here")
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text(buf.getvalue(), "cv.docx")
    assert "Tableau" in text
    assert "skills" in segment_sections(text)


def test_inline_heading():
    s = segment_sections("John Smith\nSkills: Python, Go, SQL\nExperience\nEngineer at X")
    assert "Python" in s["skills"] and "experience" in s
